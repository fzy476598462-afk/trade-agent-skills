"""Create a fictional RFQ with inline/shared text, a formula and an embedded image."""

from __future__ import annotations

import argparse
import base64
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
)
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DRAWING = """<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<xdr:oneCellAnchor><xdr:from><xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="190500" cy="190500"/>
<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1" name="Synthetic pixel"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr><a:prstGeom prst="rect"/></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>"""


def relationships(target: str, kind: str) -> bytes:
    return (
        f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/{kind}" Target="{target}"/></Relationships>'
    ).encode()


def build_demo(destination: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RFQ"
    sheet.append(["Descripción", "Part No.", "Cantidad", "Unit price", "Amount"])
    sheet.append(["Junta", "DEMO-A001", 12, 4.25, "=C2*D2"])
    sheet.append(["Soporte & montaje", "DEMO-B002", 8, 11.38, "=C3*D3"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["B"].width = 22
    for col in "CDE":
        sheet.column_dimensions[col].width = 16
    workbook.create_sheet("Notes").append(["Solicitud de cotización"])
    memory = io.BytesIO()
    workbook.save(memory)
    with zipfile.ZipFile(memory) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = (
        members["xl/worksheets/sheet1.xml"]
        .replace(
            b'<c r="A2" t="inlineStr"><is><t>Junta</t></is></c>',
            b'<c r="A2" t="s"><v>0</v></c>',
        )
        .replace(b"</worksheet>", f'<drawing xmlns:r="{REL}" r:id="rId1"/></worksheet>'.encode())
    )
    members["xl/sharedStrings.xml"] = (
        f'<sst xmlns="{NS}" count="1" uniqueCount="1"><si><r><rPr><b/></rPr><t>Junta</t></r></si></sst>'.encode()
    )
    members["xl/_rels/workbook.xml.rels"] = members["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        f'<Relationship Id="rIdDemoStrings" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'.encode(),
    )
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Default Extension="png" ContentType="image/png"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/><Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
    )
    members["xl/worksheets/_rels/sheet1.xml.rels"] = relationships(
        "../drawings/drawing1.xml", "drawing"
    )
    members["xl/drawings/drawing1.xml"] = DRAWING.encode()
    members["xl/drawings/_rels/drawing1.xml.rels"] = relationships("../media/image1.png", "image")
    members["xl/media/image1.png"] = PNG
    destination.parent.mkdir(parents=True, exist_ok=True)
    with (
        destination.open("xb") as stream,
        zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive,
    ):
        for name, content in members.items():
            archive.writestr(name, content)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_demo(args.output)
