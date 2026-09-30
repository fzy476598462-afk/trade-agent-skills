"""Read-only quote checks using explicit table and total-cell mappings."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

CENT = Decimal("0.01")
SENSITIVE = re.compile(
    r"\b(?:cost|margin|markup|profit|supplier_cost|exchange_rate)\b|成本|采购价|毛利|利润|汇率",
    re.IGNORECASE,
)
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}


def decimal(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("Expected numeric cell")
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Expected numeric cell") from exc
    if not number.is_finite() or not 0 <= number <= Decimal("1e20"):
        raise ValueError("Expected finite nonnegative number")
    return number


def mapped_rows(workbook, layout):
    tables = layout.get("tables") if isinstance(layout, dict) else None
    if not isinstance(tables, list) or not tables:
        raise ValueError("Layout requires nonempty tables")
    seen, rows = set(), []
    for table in tables:
        if not isinstance(table, dict) or set(table) - {
            "sheet",
            "start_row",
            "end_row",
            "id",
            "quality",
            "quantity",
            "unit_price",
            "amount",
        }:
            raise ValueError("Invalid table mapping")
        sheet = workbook[table["sheet"]]
        start, end = table["start_row"], table["end_row"]
        if (
            isinstance(start, bool)
            or isinstance(end, bool)
            or not isinstance(start, int)
            or not isinstance(end, int)
            or not 1 <= start <= end <= min(sheet.max_row, 100000)
        ):
            raise ValueError("Invalid table row bounds")
        for field in ("id", "quantity", "unit_price", "amount"):
            if not isinstance(table.get(field), str) or not re.fullmatch(
                r"[A-Z]{1,3}", table[field]
            ):
                raise ValueError(f"Invalid {field} column")
        if "quality" in table and (
            not isinstance(table["quality"], str)
            or not re.fullmatch(r"[A-Z]{1,3}", table["quality"])
        ):
            raise ValueError("Invalid quality column")
        for row in range(start, end + 1):
            identifier = sheet[f"{table['id']}{row}"].value
            quality = sheet[f"{table['quality']}{row}"].value if "quality" in table else ""
            if identifier is None:
                raise ValueError("Mapped rows must have an item ID")
            key = (sheet.title, str(identifier), str(quality))
            if key in seen:
                raise ValueError("Duplicate row identity; supply a quality column for options")
            seen.add(key)
            rows.append((key, sheet, row, table))
    return rows


def audit(path, layout, previous=None):
    if not isinstance(layout, dict) or set(layout) - {"tables", "totals"}:
        raise ValueError("Invalid layout fields")
    errors, warnings, prices = [], [], {}
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        for sheet in formulas:
            if sheet.sheet_state != "visible":
                errors.append({"sheet": sheet.title, "kind": "hidden_sheet"})
            for column, dimension in sheet.column_dimensions.items():
                if dimension.hidden:
                    errors.append({"sheet": sheet.title, "cell": column, "kind": "hidden_column"})
            for row, dimension in sheet.row_dimensions.items():
                if dimension.hidden:
                    errors.append({"sheet": sheet.title, "cell": str(row), "kind": "hidden_row"})
            for row in sheet:
                for cell in row:
                    if (
                        cell.data_type == "e"
                        or isinstance(cell.value, str)
                        and cell.value in ERRORS
                    ):
                        errors.append(
                            {"sheet": sheet.title, "cell": cell.coordinate, "kind": "excel_error"}
                        )
                    elif isinstance(cell.value, str) and cell.data_type != "f":
                        if SENSITIVE.search(cell.value):
                            errors.append(
                                {
                                    "sheet": sheet.title,
                                    "cell": cell.coordinate,
                                    "kind": "possible_internal_information",
                                }
                            )
                        if re.search(r"\b(?:TBD|TODO)\b", cell.value, re.IGNORECASE):
                            errors.append(
                                {
                                    "sheet": sheet.title,
                                    "cell": cell.coordinate,
                                    "kind": "unfinished_text",
                                }
                            )
                    if cell.data_type == "f" and (
                        "[" in cell.value
                        or re.search(r"\b(?:WEBSERVICE|HYPERLINK)\s*\(", cell.value, re.IGNORECASE)
                    ):
                        warnings.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "kind": "review_external_or_structured_formula",
                            }
                        )
        rows = mapped_rows(values, layout)
        for key, sheet, row, mapping in rows:
            location = {"sheet": sheet.title, "cell": f"{mapping['amount']}{row}"}
            try:
                qty = decimal(sheet[f"{mapping['quantity']}{row}"].value)
                price = decimal(sheet[f"{mapping['unit_price']}{row}"].value)
                amount = decimal(sheet[f"{mapping['amount']}{row}"].value)
            except ValueError:
                errors.append(
                    {**location, "kind": "missing_or_invalid_numeric_value_or_formula_cache"}
                )
                continue
            prices[key] = price
            if price != price.quantize(CENT, rounding=ROUND_HALF_UP):
                errors.append({**location, "kind": "unit_price_exceeds_two_decimal_contract"})
            expected = (qty * price).quantize(CENT, rounding=ROUND_HALF_UP)
            if abs(expected - amount) > Decimal("0.005"):
                errors.append(
                    {
                        **location,
                        "kind": "line_amount_mismatch",
                        "expected": str(expected),
                        "actual": str(amount),
                    }
                )
        totals = layout.get("totals", [])
        if not isinstance(totals, list):
            raise ValueError("totals must be a list")
        for total in totals:
            if (
                not isinstance(total, dict)
                or set(total) != {"sheet", "cell", "sources"}
                or not isinstance(total["sources"], list)
                or not total["sources"]
            ):
                raise ValueError("Invalid total mapping")
            if not isinstance(total["cell"], str) or not re.fullmatch(
                r"[A-Z]{1,3}[1-9][0-9]{0,6}", total["cell"]
            ):
                raise ValueError("Total cell must be a single coordinate")
            actual_cell = values[total["sheet"]][total["cell"]]
            expected = Decimal(0)
            try:
                for source in total["sources"]:
                    if not isinstance(source, list) or len(source) != 2:
                        raise ValueError("Total sources must be [sheet, range] pairs")
                    min_col, min_row, max_col, max_row = range_boundaries(source[1])
                    if None in (min_col, min_row, max_col, max_row) or max_row * max_col > 1000000:
                        raise ValueError("Invalid source range")
                    for row in values[source[0]].iter_rows(
                        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
                    ):
                        for cell in row:
                            expected += decimal(cell.value)
                actual = decimal(actual_cell.value)
                if abs(expected - actual) > Decimal("0.005"):
                    errors.append(
                        {
                            "sheet": total["sheet"],
                            "cell": total["cell"],
                            "kind": "total_mismatch",
                            "expected": str(expected),
                            "actual": str(actual),
                        }
                    )
            except ValueError:
                errors.append(
                    {
                        "sheet": total["sheet"],
                        "cell": total["cell"],
                        "kind": "missing_or_invalid_total_value_or_cache",
                    }
                )
        changes = []
        if previous:
            old = load_workbook(previous, data_only=True)
            try:
                for key, sheet, row, mapping in mapped_rows(old, layout):
                    prior = decimal(sheet[f"{mapping['unit_price']}{row}"].value)
                    if key in prices and prices[key] != prior:
                        changes.append(
                            {
                                "sheet": key[0],
                                "item_id": key[1],
                                "quality": key[2],
                                "old": str(prior),
                                "new": str(prices[key]),
                            }
                        )
            finally:
                old.close()
        return {
            "ok": not errors,
            "checked_lines": len(rows),
            "errors": errors,
            "warnings": warnings,
            "price_changes": changes,
            "coverage": "Only explicitly mapped rows and totals are checked arithmetically; text/hidden-content checks cover all loaded sheets. Not a complete confidential-data detector.",
        }
    finally:
        values.close()
        formulas.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("layout", type=Path)
    parser.add_argument("--previous", type=Path)
    args = parser.parse_args()
    try:
        result = audit(
            args.workbook, json.loads(args.layout.read_text(encoding="utf-8-sig")), args.previous
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if not result["ok"]:
            parser.exit(1)
    except (ValueError, KeyError, TypeError, OSError, InvalidOperation, zipfile.BadZipFile) as exc:
        parser.exit(2, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
