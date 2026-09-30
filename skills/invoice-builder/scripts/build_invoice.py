"""Build new trade documents from explicit data; never edit an existing workbook."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, DecimalException, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CENT = Decimal("0.01")
ACCENT = "0F766E"
INK = "172B36"
TITLES = {
    "commercial": "COMMERCIAL INVOICE",
    "proforma": "PROFORMA INVOICE",
    "customs": "CUSTOMS INVOICE",
    "packing": "PACKING LIST",
}
HEADERS = {
    "commercial": [
        "NO.",
        "DESCRIPTION",
        "PART NO.",
        "MODEL",
        "QTY",
        "UNIT",
        "UNIT PRICE",
        "AMOUNT",
    ],
    "proforma": ["NO.", "DESCRIPTION", "PART NO.", "MODEL", "QTY", "UNIT", "UNIT PRICE", "AMOUNT"],
    "customs": ["NO.", "DESCRIPTION", "MATERIAL", "USE", "HS CODE", "QTY", "UNIT PRICE", "AMOUNT"],
    "packing": ["NO.", "DESCRIPTION", "PART NO.", "QTY", "CARTONS", "NET KG", "GROSS KG", "CBM"],
}


def number(value: object, field: str, *, positive: bool = False) -> Decimal:
    """Parse finite, nonnegative decimal input without accepting booleans."""
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{field}: expected a number")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"{field}: invalid number") from exc
    if not result.is_finite() or result < 0 or (positive and result == 0):
        raise ValueError(f"{field}: must be finite and {'positive' if positive else 'nonnegative'}")
    if result > Decimal("1e9") or (result and result < Decimal("1e-9")):
        raise ValueError(f"{field}: value is outside the supported range (1e-9 to 1e9, or zero)")
    return result


def text(value: object, field: str, *, required: bool = False) -> str:
    if not isinstance(value, str) or (required and not value.strip()):
        raise ValueError(f"{field}: expected {'nonempty ' if required else ''}text")
    if len(value) > 2000 or any(ord(char) < 32 and char not in "\n\t" for char in value):
        raise ValueError(f"{field}: text is too long or contains control characters")
    return value.strip()


def keys(data: object, allowed: set[str], field: str) -> dict:
    if not isinstance(data, dict):
        raise ValueError(f"{field}: expected an object")
    extra = set(data) - allowed
    if extra:
        raise ValueError(f"{field}: unsupported fields {', '.join(sorted(extra))}")
    return data


def validate(data: dict, kind: str) -> tuple[dict, list[dict], Decimal, Decimal]:
    """Normalize all inputs before creating any output file."""
    keys(
        data,
        {"invoice_no", "date", "currency", "seller", "buyer", "items", "freight", "terms"},
        "document",
    )
    data = dict(data)
    data["invoice_no"] = text(data.get("invoice_no"), "invoice_no", required=True)
    data["date"] = text(data.get("date"), "date", required=True)
    date.fromisoformat(data["date"])
    data["currency"] = data.get("currency", "USD")
    if not isinstance(data["currency"], str) or not re.fullmatch(r"[A-Z]{3}", data["currency"]):
        raise ValueError("currency: use a three-letter uppercase currency code")
    for role in ("seller", "buyer"):
        party = keys(data.get(role), {"name", "address", "contact"}, role)
        data[role] = {
            key: text(party.get(key, ""), f"{role}.{key}", required=key == "name")
            for key in ("name", "address", "contact")
        }
    raw_items = data.get("items")
    if not isinstance(raw_items, list) or not raw_items or len(raw_items) > 200:
        raise ValueError("items: supply 1 to 200 line items")
    items = []
    subtotal = Decimal(0)
    for index, raw in enumerate(raw_items, 1):
        prefix = f"items[{index}]"
        allowed = {"description", "quantity", "part_no", "unit"}
        if kind == "packing":
            allowed |= {"cartons", "net_kg", "gross_kg", "cbm"}
        else:
            allowed |= {"unit_price", "model", "material", "use", "hs_code"}
        keys(raw, allowed, prefix)
        item = {"description": text(raw.get("description"), prefix, required=True)}
        item["quantity"] = number(raw.get("quantity"), f"{prefix}.quantity", positive=True)
        for key in allowed - {
            "description",
            "quantity",
            "unit_price",
            "cartons",
            "net_kg",
            "gross_kg",
            "cbm",
        }:
            item[key] = text(raw.get(key, "pcs" if key == "unit" else ""), f"{prefix}.{key}")
        if kind == "packing":
            for key in ("cartons", "net_kg", "gross_kg", "cbm"):
                item[key] = number(raw[key], f"{prefix}.{key}") if key in raw else None
            if item["cartons"] is not None and (
                item["cartons"] < 1 or item["cartons"] != item["cartons"].to_integral_value()
            ):
                raise ValueError(f"{prefix}.cartons: must be a positive integer")
            if item["net_kg"] is not None and item["gross_kg"] is not None:
                if item["net_kg"] > item["gross_kg"]:
                    raise ValueError(f"{prefix}: net weight exceeds gross weight")
        else:
            item["unit_price"] = number(raw.get("unit_price"), f"{prefix}.unit_price").quantize(
                CENT, rounding=ROUND_HALF_UP
            )
            item["amount"] = (item["quantity"] * item["unit_price"]).quantize(
                CENT, rounding=ROUND_HALF_UP
            )
            subtotal += item["amount"]
        items.append(item)
    if kind in {"customs", "packing"} and "freight" in data:
        raise ValueError(f"{kind}: freight does not belong in this document")
    freight = number(data.get("freight", 0), "freight").quantize(CENT, rounding=ROUND_HALF_UP)
    data["terms"] = text(data.get("terms", ""), "terms")
    return data, items, subtotal, freight


def cell(ws, row: int, column: int, value: object, *, bold: bool = False, color: str = INK):
    result = ws.cell(row, column)
    if isinstance(value, str):
        result.value = value
        result.data_type = "s"  # Never interpret user descriptions as Excel formulas.
    elif isinstance(value, Decimal):
        result.value = float(value)
    else:
        result.value = value
    result.font = Font(name="Calibri", size=11, bold=bold, color=color)
    result.alignment = Alignment(vertical="center", wrap_text=True)
    return result


def merged(ws, row: int, start: int, end: int, value: str, **style):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    return cell(ws, row, start, value, **style)


def build_document(data: dict, output: Path, kind: str = "commercial") -> dict:
    if kind not in TITLES:
        raise ValueError(f"Unknown document kind: {kind}")
    data, items, subtotal, freight = validate(data, kind)
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have the .xlsx extension")
    wb = Workbook()
    ws = wb.active
    ws.title = TITLES[kind].title()
    ws.sheet_view.showGridLines = False
    for col, width in enumerate((6, 42, 20, 22, 16, 14, 19, 19), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    merged(ws, 1, 1, 8, TITLES[kind], bold=True, color=ACCENT).font = Font(
        name="Calibri", size=22, bold=True, color=ACCENT
    )
    merged(
        ws, 2, 1, 8, f"No. {data['invoice_no']}    |    {data['date']}    |    {data['currency']}"
    )
    ws.row_dimensions[1].height = 36
    for role, start in (("seller", 4), ("buyer", 8)):
        party = data[role]
        merged(ws, start, 1, 8, f"{role.upper()}  {party['name']}", bold=True)
        merged(ws, start + 1, 1, 8, party["address"])
        merged(ws, start + 2, 1, 8, party["contact"])
        for row in range(start, start + 3):
            ws.row_dimensions[row].height = 32 if row > start else 24
    header = 12
    for col, label in enumerate(HEADERS[kind], 1):
        item_cell = cell(ws, header, col, label, bold=True, color="FFFFFF")
        item_cell.fill = PatternFill("solid", fgColor=ACCENT)
    ws.row_dimensions[header].height = 32
    border = Border(bottom=Side(style="thin", color="DDE7E6"))
    for index, item in enumerate(items, 1):
        row = header + index
        if kind == "packing":
            values = [
                index,
                item["description"],
                item["part_no"],
                item["quantity"],
                item["cartons"],
                item["net_kg"],
                item["gross_kg"],
                item["cbm"],
            ]
        elif kind == "customs":
            values = [
                index,
                item["description"],
                item["material"],
                item["use"],
                item["hs_code"],
                item["quantity"],
                item["unit_price"],
                item["amount"],
            ]
        else:
            values = [
                index,
                item["description"],
                item["part_no"],
                item["model"],
                item["quantity"],
                item["unit"],
                item["unit_price"],
                item["amount"],
            ]
        for col, value in enumerate(values, 1):
            item_cell = cell(ws, row, col, value)
            item_cell.border = border
            if index % 2 == 0:
                item_cell.fill = PatternFill("solid", fgColor="F2F7F6")
            if isinstance(value, Decimal):
                item_cell.number_format = (
                    "#,##0.00" if col in (7, 8) and kind != "packing" else "0.###"
                )
        ws.row_dimensions[row].height = max(40, 16 * (len(item["description"]) // 35 + 1))
    last = header + len(items) + 2
    if kind != "packing":
        totals = [("Subtotal (goods)", subtotal)]
        if kind in {"commercial", "proforma"}:
            if "freight" in data:
                totals.append(("Freight", freight))
            totals.append((f"TOTAL ({data['currency']})", subtotal + freight))
        for label, value in totals:
            merged(ws, last, 5, 7, label, bold=True)
            cell(ws, last, 8, value, bold=True, color=ACCENT).number_format = "#,##0.00"
            ws.row_dimensions[last].height = 28
            last += 1
    if data["terms"]:
        merged(ws, last + 1, 1, 8, data["terms"])
        ws.row_dimensions[last + 1].height = max(32, 16 * (len(data["terms"]) // 100 + 1))
        last += 2
    ws.freeze_panes = "C13"
    ws.print_title_rows = "1:12"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:H{last + 1}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as stream:
        try:
            wb.save(stream)
        except Exception:
            stream.close()
            output.unlink()
            raise
    return {
        "output": str(output),
        "kind": kind,
        "lines": len(items),
        "goods": str(subtotal),
        "total": str(subtotal + freight) if kind in {"commercial", "proforma"} else None,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="JSON document data")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--kind", choices=TITLES, default="commercial")
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
        print(json.dumps(build_document(data, args.output, args.kind), ensure_ascii=False))
    except (ValueError, OSError, TypeError, KeyError, DecimalException) as exc:
        parser.exit(2, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
