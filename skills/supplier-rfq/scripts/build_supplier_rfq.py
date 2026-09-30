"""Create a supplier RFQ with reference images embedded beside their line items."""

from __future__ import annotations

import argparse
import io
import json
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

ACCENT = "0F766E"


def text(value, name, required=False):
    if not isinstance(value, str) or (required and not value.strip()):
        raise ValueError(f"{name}: expected {'nonempty ' if required else ''}text")
    if len(value) > 1500 or any(ord(c) < 32 and c not in "\n\t" for c in value):
        raise ValueError(f"{name}: unsupported text")
    return value.strip()


def allowed(data, fields, name):
    if not isinstance(data, dict) or set(data) - fields:
        raise ValueError(f"{name}: invalid object or unsupported fields")


def quantity(value):
    if isinstance(value, bool):
        raise ValueError("quantity: expected a positive number")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("quantity: invalid number") from exc
    if not result.is_finite() or not 0 < result <= 1000000:
        raise ValueError("quantity: expected a positive finite number <= 1000000")
    return float(result)


def picture(path):
    with Image.open(path) as source:
        source.load()
        image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail((280, 160))
        stream = io.BytesIO()
        image.save(stream, format="PNG")  # Re-encode without original EXIF/GPS metadata.
        stream.seek(0)
        return stream


def build(data, output, image_root=Path(".")):
    allowed(data, {"rfq_no", "date", "vehicle", "instructions", "items"}, "RFQ")
    number = text(data.get("rfq_no"), "rfq_no", True)
    dated = text(data.get("date"), "date", True)
    date.fromisoformat(dated)
    vehicle = text(data.get("vehicle", ""), "vehicle")
    instructions = text(data.get("instructions", ""), "instructions")
    raw = data.get("items")
    if not isinstance(raw, list) or not 1 <= len(raw) <= 200:
        raise ValueError("items: supply 1 to 200 lines")
    rows, images = [], []
    for item in raw:
        allowed(
            item, {"description", "source_text", "part_no", "quantity", "photo", "notes"}, "item"
        )
        row = {
            key: text(item.get(key, ""), key, key == "description")
            for key in ("description", "source_text", "part_no", "notes")
        }
        row["quantity"] = quantity(item.get("quantity"))
        rows.append(row)
        photo = text(item.get("photo", ""), "photo")
        images.append(picture(Path(image_root) / photo) if photo else None)
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be XLSX")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Supplier RFQ"
    sheet.sheet_view.showGridLines = False

    def cell(row, col, value, bold=False):
        result = sheet.cell(row, col, value)
        if isinstance(value, str):
            result.data_type = "s"
        result.font = Font(name="Calibri", size=11, bold=bold, color="172B36")
        result.alignment = Alignment(vertical="center", wrap_text=True)
        return result

    for row, value in enumerate(
        ("SUPPLIER REQUEST FOR QUOTATION", f"{number} | {dated}", vehicle, instructions), 1
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell(row, 1, value, row == 1)
        sheet.row_dimensions[row].height = 32
    labels = [
        "NO.",
        "DESCRIPTION",
        "SOURCE TEXT",
        "PART NO.",
        "QTY",
        "REFERENCE IMAGE",
        "QUOTE A",
        "QUOTE B",
        "LEAD TIME",
        "NOTES",
    ]
    for col, label in enumerate(labels, 1):
        head = cell(6, col, label, True)
        head.fill = PatternFill("solid", fgColor=ACCENT)
        head.font = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
    sheet.row_dimensions[6].height = 28
    for index, (item, stream) in enumerate(zip(rows, images, strict=True), 1):
        row = 6 + index
        values = [
            index,
            item["description"],
            item["source_text"],
            item["part_no"],
            item["quantity"],
            "",
            "",
            "",
            "",
            item["notes"],
        ]
        for col, value in enumerate(values, 1):
            entry = cell(row, col, value)
            if index % 2 == 0:
                entry.fill = PatternFill("solid", fgColor="F2F7F6")
        sheet.row_dimensions[row].height = 128 if stream else 44
        if stream:
            sheet.add_image(ExcelImage(stream), f"F{row}")
    for col, width in enumerate((6, 24, 24, 20, 8, 41, 14, 14, 16, 28), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "F7"
    sheet.print_title_rows = "1:6"
    sheet.print_area = f"A1:J{6 + len(rows)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        with output.open("xb") as destination:
            try:
                workbook.save(destination)
            except Exception:
                destination.close()
                output.unlink()
                raise
    finally:
        for stream in images:
            if stream and not stream.closed:
                stream.close()
    with zipfile.ZipFile(output) as archive:
        count = sum(name.startswith("xl/media/") for name in archive.namelist())
    if count != sum(stream is not None for stream in images):
        raise ValueError("Embedded image verification failed")
    return {"output": str(output), "lines": len(rows), "images": count}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = build(
            json.loads(args.input.read_text(encoding="utf-8-sig")), args.output, args.input.parent
        )
        print(json.dumps(result, ensure_ascii=False))
    except (ValueError, TypeError, OSError) as exc:
        parser.exit(2, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
