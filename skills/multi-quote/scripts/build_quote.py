"""Build a multi-vehicle, multi-quality quote with an editable quantity per option."""

from __future__ import annotations

import argparse
import copy
import io
import json
import re
import zipfile
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
CENT = Decimal("0.01")


def number(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name}: expected number")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"{name}: invalid number") from exc
    if not result.is_finite() or not 0 <= result <= 1000000000:
        raise ValueError(f"{name}: expected finite nonnegative number <= 1e9")
    return result


def text(value, name):
    if not isinstance(value, str) or not value.strip() or len(value) > 1000:
        raise ValueError(f"{name}: supply nonempty text <= 1000 characters")
    if any(ord(c) < 32 and c not in "\n\t" for c in value):
        raise ValueError(f"{name}: invalid text")
    return value.strip()


def fields(value, allowed, label):
    if not isinstance(value, dict) or set(value) - allowed:
        raise ValueError(f"{label}: invalid object or unsupported fields")


def build(data, output):
    fields(
        data,
        {"quote_no", "date", "currency", "seller", "buyer", "vehicles", "freight", "terms"},
        "quote",
    )
    quote_no, dated = text(data.get("quote_no"), "quote_no"), text(data.get("date"), "date")
    date.fromisoformat(dated)
    currency = data.get("currency")
    if not isinstance(currency, str) or not re.fullmatch(r"[A-Z]{3}", currency):
        raise ValueError("currency: supply uppercase three-letter code")
    seller, buyer = text(data.get("seller"), "seller"), text(data.get("buyer"), "buyer")
    terms = text(data["terms"], "terms") if data.get("terms") else ""
    freight = number(data.get("freight", 0), "freight").quantize(CENT, rounding=ROUND_HALF_UP)
    vehicles = data.get("vehicles")
    if not isinstance(vehicles, list) or not 1 <= len(vehicles) <= 30:
        raise ValueError("vehicles: supply 1 to 30 entries")
    normalized, identifiers = [], set()
    for vehicle in vehicles:
        fields(vehicle, {"sheet", "model", "items"}, "vehicle")
        name, model = text(vehicle.get("sheet"), "sheet"), text(vehicle.get("model"), "model")
        if (
            len(name) > 31
            or re.search(r"[\[\]:*?/\\]", name)
            or name.casefold() in identifiers | {"summary"}
        ):
            raise ValueError("Sheet names must be valid, unique and different from Summary")
        identifiers.add(name.casefold())
        items = vehicle.get("items")
        if not isinstance(items, list) or not 1 <= len(items) <= 200:
            raise ValueError("Each vehicle needs 1 to 200 items")
        rows, part_ids = [], set()
        for item in items:
            fields(item, {"id", "description", "part_no", "options", "notes"}, "item")
            item_id = text(item.get("id"), "item.id")
            if item_id in part_ids:
                raise ValueError("Item IDs must be unique within each vehicle")
            part_ids.add(item_id)
            description = text(item.get("description"), "description")
            part_no = text(item["part_no"], "part_no") if item.get("part_no") else ""
            notes = text(item["notes"], "notes") if item.get("notes") else ""
            options = item.get("options")
            if not isinstance(options, list) or not 1 <= len(options) <= 4:
                raise ValueError("Each item needs 1 to 4 quality/price options")
            qualities = set()
            for option in options:
                fields(option, {"quality", "quantity", "unit_price"}, "option")
                quality = text(option.get("quality"), "quality")
                if quality.casefold() in qualities:
                    raise ValueError("Quality option labels must be unique within an item")
                qualities.add(quality.casefold())
                qty = number(option.get("quantity", 0), "quantity")
                price = number(option.get("unit_price"), "unit_price").quantize(
                    CENT, rounding=ROUND_HALF_UP
                )
                amount = (qty * price).quantize(CENT, rounding=ROUND_HALF_UP)
                rows.append([item_id, description, part_no, quality, qty, price, amount, notes])
        normalized.append((name, model, rows))
    workbook = Workbook()
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    summary = workbook.active
    summary.title = "Summary"
    caches, totals = {}, []

    def cell(sheet, row, col, value, bold=False):
        result = sheet.cell(row, col, float(value) if isinstance(value, Decimal) else value)
        if isinstance(value, str):
            result.data_type = "s"
        result.font = Font(name="Calibri", size=11, bold=bold, color="172B36")
        result.alignment = Alignment(vertical="center", wrap_text=True)
        if isinstance(value, Decimal):
            result.number_format = "#,##0.00"
        return result

    for sheet_index, (name, model, rows) in enumerate(normalized, 2):
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for row, value in enumerate(
            (
                f"QUOTE {quote_no} | {model}",
                f"{dated} | {currency}",
                f"Seller: {seller} | Buyer: {buyer}",
            ),
            1,
        ):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            cell(sheet, row, 1, value, row == 1)
            sheet.row_dimensions[row].height = 30
        for col, label in enumerate(
            (
                "ITEM ID",
                "DESCRIPTION",
                "PART NO.",
                "QUALITY",
                "QTY",
                "UNIT PRICE",
                "AMOUNT",
                "NOTES",
            ),
            1,
        ):
            head = cell(sheet, 5, col, label, True)
            head.fill = PatternFill("solid", fgColor="0F766E")
            head.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cached = {}
        for index, values in enumerate(rows, 6):
            for col, value in enumerate(values, 1):
                entry = cell(sheet, index, col, value)
                if index % 2 == 0:
                    entry.fill = PatternFill("solid", fgColor="F2F7F6")
            sheet.cell(index, 7, f"=ROUND(E{index}*F{index},2)")
            cached[f"G{index}"] = values[6]
            sheet.row_dimensions[index].height = max(34, 16 * (len(values[1]) // 35 + 1))
        subtotal = sum((row[6] for row in rows), Decimal(0))
        last = len(rows) + 7
        cell(sheet, last, 6, "SUBTOTAL", True)
        cell(sheet, last, 7, subtotal, True)
        sheet.cell(last, 7, f"=SUM(G6:G{len(rows) + 5})")
        cached[f"G{last}"] = subtotal
        caches[f"xl/worksheets/sheet{sheet_index}.xml"] = cached
        totals.append((name, last, subtotal))
        for col, width in enumerate((16, 40, 22, 22, 12, 18, 18, 28), 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "E6"
        sheet.print_title_rows = "1:5"
        sheet.print_area = f"A1:H{last}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:D1")
    cell(summary, 1, 1, f"QUOTE {quote_no}", True)
    summary.merge_cells("A2:D2")
    cell(summary, 2, 1, f"{dated} | {currency} | {seller} → {buyer}")
    cell(summary, 4, 1, "VEHICLE", True)
    cell(summary, 4, 4, "SELECTED GOODS", True)
    cached = {}
    for row, (name, last, subtotal) in enumerate(totals, 5):
        cell(summary, row, 1, name)
        summary.cell(row, 4, f"='{name.replace(chr(39), chr(39) * 2)}'!G{last}")
        summary.cell(row, 4).number_format = "#,##0.00"
        cached[f"D{row}"] = subtotal
    goods = sum((total[2] for total in totals), Decimal(0))
    row = len(totals) + 6
    for label, value in (
        ("Goods subtotal", goods),
        ("Freight", freight),
        (f"TOTAL ({currency})", goods + freight),
    ):
        cell(summary, row, 1, label, True)
        cell(summary, row, 4, value, True)
        if label == "Goods subtotal":
            summary.cell(row, 4, f"=SUM(D5:D{len(totals) + 4})")
            goods_row = row
            cached[f"D{row}"] = value
        elif label.startswith("TOTAL"):
            summary.cell(row, 4, f"=D{goods_row}+D{goods_row + 1}")
            cached[f"D{row}"] = value
        row += 1
    summary.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=4)
    cell(
        summary,
        row + 1,
        1,
        terms or "Quantities select options. Confirm the final selection before ordering.",
    )
    summary.row_dimensions[row + 1].height = 48
    for col, width in enumerate((32, 24, 20, 24), 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.row_dimensions[1].height = 34
    summary.row_dimensions[2].height = 40
    summary.print_area = f"A1:D{row + 1}"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    caches["xl/worksheets/sheet1.xml"] = cached
    memory = io.BytesIO()
    workbook.save(memory)
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be XLSX")
    output.parent.mkdir(parents=True, exist_ok=True)
    finished = io.BytesIO()
    with zipfile.ZipFile(memory) as source, zipfile.ZipFile(finished, "w") as destination:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename in caches:
                root = ET.fromstring(payload)
                for node in root.iter(f"{{{NS}}}c"):
                    if node.get("r") in caches[info.filename]:
                        value = node.find(f"{{{NS}}}v")
                        if value is None:
                            value = ET.SubElement(node, f"{{{NS}}}v")
                        value.text = str(caches[info.filename][node.get("r")])
                payload = ET.tostring(root, encoding="utf-8")
            destination.writestr(copy.copy(info), payload)
    with output.open("xb") as stream:
        try:
            stream.write(finished.getvalue())
        except OSError:
            stream.close()
            output.unlink()
            raise
    return {
        "output": str(output),
        "vehicles": len(totals),
        "selected_goods": str(goods),
        "total": str(goods + freight),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        print(
            json.dumps(build(json.loads(args.input.read_text(encoding="utf-8-sig")), args.output))
        )
    except (ValueError, TypeError, OSError, InvalidOperation) as exc:
        parser.exit(2, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
