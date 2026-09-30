"""Public trade adaptations: arithmetic, privacy and package-preservation checks."""

import copy
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]


def module(skill, script):
    spec = importlib.util.spec_from_file_location(
        skill.replace("-", "_"), ROOT / "skills" / skill / "scripts" / script
    )
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


rfq = module("supplier-rfq", "build_supplier_rfq.py")
quote = module("multi-quote", "build_quote.py")
audit = module("quote-audit", "audit_quote.py")
edit = module("xlsx-edit", "edit_xlsx.py")
landed = module("landed-cost", "compare_costs.py")


def sample(name):
    return json.loads((ROOT / "examples" / name).read_text(encoding="utf-8"))


class NewTools(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def workbook(self):
        path = self.directory / "quote.xlsx"
        quote.build(sample("multi-quote.json"), path)
        return path

    def test_supplier_images_quantities_and_literal_text(self):
        picture = self.directory / "part.jpg"
        image = Image.new("RGB", (400, 200), "teal")
        exif = Image.Exif()
        exif[315] = "PRIVATE CAMERA OWNER"
        image.save(picture, exif=exif)
        data = sample("supplier-rfq.json")
        data["items"][0]["photo"] = "part.jpg"
        data["items"][0]["source_text"] = '=HYPERLINK("http://example.com")'
        output = self.directory / "rfq.xlsx"
        result = rfq.build(data, output, self.directory)
        self.assertEqual(result["images"], 1)
        book = load_workbook(output)
        self.assertEqual(book.active["E7"].value, 4)
        self.assertEqual(book.active["C7"].data_type, "s")
        self.assertEqual(book.active._images[0].anchor._from.row, 6)
        book.close()
        with zipfile.ZipFile(output) as archive:
            pixels = archive.read("xl/media/image1.png")
            self.assertNotIn(b"PRIVATE CAMERA OWNER", pixels)
            self.assertNotIn(b"eXIf", pixels)
        with self.assertRaises(FileExistsError):
            rfq.build(data, output, self.directory)

    def test_supplier_rejects_invalid_qty_unknown_fields_and_missing_image(self):
        for qty in (0, -1, True, "NaN"):
            data = sample("supplier-rfq.json")
            data["items"][0]["quantity"] = qty
            with self.assertRaises(ValueError):
                rfq.build(data, self.directory / "no.xlsx")
        data = sample("supplier-rfq.json")
        data["customer_phone"] = "must not leak"
        with self.assertRaises(ValueError):
            rfq.build(data, self.directory / "no.xlsx")
        data.pop("customer_phone")
        with self.assertRaises(FileNotFoundError):
            rfq.build(data, self.directory / "no.xlsx", self.directory)

    def test_quote_formula_caches_and_audit_are_consistent(self):
        path = self.workbook()
        values, formulas = load_workbook(path, data_only=True), load_workbook(path)
        self.assertEqual(values["Summary"]["D10"].value, 155)
        self.assertEqual(values["Compact"]["G7"].value, 0)
        self.assertEqual(formulas["Compact"]["G6"].value, "=ROUND(E6*F6,2)")
        self.assertTrue(formulas.calculation.fullCalcOnLoad)
        self.assertTrue(audit.audit(path, sample("quote-layout.json"))["ok"])
        values.close()
        formulas.close()

    def test_quote_options_stay_separate_and_invalid_inputs_fail(self):
        data = sample("multi-quote.json")
        data["vehicles"][0]["items"][0]["options"][1]["quantity"] = 2
        output = self.directory / "q.xlsx"
        self.assertEqual(quote.build(data, output)["total"], "171.50")
        with self.assertRaises(FileExistsError):
            quote.build(data, output)
        for mutate in (
            lambda d: d["vehicles"][1].update(sheet="Compact"),
            lambda d: d["vehicles"][0]["items"][0]["options"][0].update(unit_price=-1),
            lambda d: d.update(internal_margin=20),
        ):
            data = sample("multi-quote.json")
            mutate(data)
            with self.assertRaises(ValueError):
                quote.build(data, self.directory / "invalid.xlsx")

    def test_audit_detects_arithmetic_hidden_columns_and_sensitive_labels(self):
        path = self.workbook()
        book = load_workbook(path)
        book["Compact"]["G6"] = 999
        book["Compact"]["K5"] = "Internal margin"
        book["Compact"].column_dimensions["K"].hidden = True
        book.save(path)
        book.close()
        result = audit.audit(path, sample("quote-layout.json"))
        self.assertFalse(result["ok"])
        kinds = {error["kind"] for error in result["errors"]}
        self.assertIn("line_amount_mismatch", kinds)
        self.assertIn("hidden_column", kinds)
        self.assertIn("possible_internal_information", kinds)
        self.assertIn("missing_or_invalid_numeric_value_or_formula_cache", kinds)

    def test_audit_reports_price_changes_by_item_and_quality(self):
        old = self.workbook()
        data = sample("multi-quote.json")
        data["vehicles"][0]["items"][0]["options"][0]["unit_price"] = 15
        current = self.directory / "new.xlsx"
        quote.build(data, current)
        result = audit.audit(current, sample("quote-layout.json"), previous=old)
        self.assertTrue(result["ok"])
        self.assertEqual(len(result["price_changes"]), 1)
        self.assertEqual(result["price_changes"][0]["item_id"], "P-001")
        self.assertEqual(result["price_changes"][0]["old"], "12.5")

    def test_edit_preserves_images_styles_expressions_and_clears_caches(self):
        data = sample("supplier-rfq.json")
        image = self.directory / "demo-part.png"
        Image.new("RGB", (80, 40), "teal").save(image)
        source = self.directory / "rfq.xlsx"
        rfq.build(data, source, self.directory)
        # Include a dependent formula to test cache invalidation across sheets.
        book = load_workbook(source)
        book.create_sheet("Total")["A1"] = "=SUM('Supplier RFQ'!E7:E8)"
        book.save(source)
        book.close()
        patch = edit.inspect(source, [("Supplier RFQ", "E7"), ("Supplier RFQ", "C8")])
        patch["changes"][0]["value"] = 6
        patch["changes"][1]["value"] = '=TEXT("& < ")'
        output = self.directory / "edited.xlsx"
        result = edit.apply(source, patch, output)
        self.assertEqual(result["formula_caches_cleared"], 1)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
            for member in before.namelist():
                if member not in result["changed_members"]:
                    self.assertEqual(before.read(member), after.read(member), member)
        book = load_workbook(output)
        self.assertEqual(book["Supplier RFQ"]["E7"].value, 6)
        self.assertEqual(book["Supplier RFQ"]["C8"].data_type, "s")
        self.assertEqual(book["Total"]["A1"].value, "=SUM('Supplier RFQ'!E7:E8)")
        self.assertEqual(len(book["Supplier RFQ"]._images), 1)
        self.assertTrue(book.calculation.forceFullCalc)
        book.close()

    def test_edit_cannot_overwrite_edit_formula_or_use_stale_values(self):
        source = self.workbook()
        patch = edit.inspect(source, [("Compact", "E6")])
        patch["changes"][0]["value"] = 6
        original = source.read_bytes()
        for changed in (
            {**patch, "source_sha256": "0" * 64},
            {**patch, "changes": [{**patch["changes"][0], "expected": 999}]},
            {**patch, "changes": [{**patch["changes"][0], "cell": "G6"}]},
            {**patch, "changes": [{**patch["changes"][0], "cell": "Z999"}]},
            {**patch, "changes": [patch["changes"][0], patch["changes"][0]]},
        ):
            with self.assertRaises(ValueError):
                edit.apply(source, changed, self.directory / "invalid.xlsx")
        self.assertEqual(source.read_bytes(), original)
        with self.assertRaises(ValueError):
            edit.apply(source, patch, source)

    def test_edit_invalidates_every_formula_cache_and_audit_requires_recalculation(self):
        source = self.workbook()
        patch = edit.inspect(source, [("Compact", "E6")])
        patch["changes"][0]["value"] = 6
        output = self.directory / "edited.xlsx"
        result = edit.apply(source, patch, output)
        self.assertGreater(result["formula_caches_cleared"], 5)
        with zipfile.ZipFile(output) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    for cell in ET.fromstring(archive.read(name)).iter(f"{{{edit.NS}}}c"):
                        if cell.find(f"{{{edit.NS}}}f") is not None:
                            self.assertIsNone(cell.find(f"{{{edit.NS}}}v"))
        self.assertFalse(audit.audit(output, sample("quote-layout.json"))["ok"])

    def test_landed_costs_fx_ties_and_unknowns(self):
        data = sample("landed-cost.json")
        result = landed.compare(data)
        self.assertTrue(result["complete"])
        self.assertEqual(result["lowest_cost_offers"], ["Offer A"])
        self.assertEqual(result["offers"][0]["total_target"], "145.00")
        self.assertEqual(result["offers"][1]["total_target"], "154.00")
        data["offers"][1] = copy.deepcopy(data["offers"][0])
        data["offers"][1]["name"] = "Offer B"
        self.assertEqual(landed.compare(data)["lowest_cost_offers"], ["Offer A", "Offer B"])
        del data["offers"][0]["costs"]["tax"]
        result = landed.compare(data)
        self.assertFalse(result["complete"])
        self.assertEqual(result["lowest_cost_offers"], [])
        self.assertIsNone(result["offers"][0]["total_target"])

    def test_landed_missing_fx_invalid_fees_and_unconfirmed_scope(self):
        data = sample("landed-cost.json")
        del data["offers"][1]["fx_to_target"]
        self.assertIn("fx_to_target", landed.compare(data)["offers"][1]["missing"])
        data["offers"][0]["costs"]["goods"] = -1
        with self.assertRaises(ValueError):
            landed.compare(data)
        data = sample("landed-cost.json")
        data["comparison_scope"]["quantity"] = 0
        with self.assertRaises(ValueError):
            landed.compare(data)


if __name__ == "__main__":
    unittest.main()
