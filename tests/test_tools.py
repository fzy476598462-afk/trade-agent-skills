"""Regression tests for arithmetic, document boundaries and XLSX preservation."""

import copy
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def load(relative):
    spec = importlib.util.spec_from_file_location(Path(relative).stem, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


invoice = load("skills/invoice-builder/scripts/build_invoice.py")
freight = load("skills/freight-estimator/scripts/estimate_freight.py")
translate = load("skills/translate-rfq/scripts/translate_xlsx.py")
demo = load("examples/make_demo_rfq.py")


def example(name):
    return json.loads((ROOT / "examples" / name).read_text(encoding="utf-8"))


class InvoiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.output = Path(self.temp.name) / "invoice.xlsx"
        self.data = example("invoice.json")

    def test_visible_rounding_and_literal_text(self):
        self.data["items"][0]["description"] = '=HYPERLINK("https://example.com")'
        result = invoice.build_document(self.data, self.output)
        self.assertEqual(result["total"], "142.04")
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active["G14"].value, 11.38)
        self.assertEqual(wb.active["H14"].value, 91.04)
        self.assertEqual(wb.active["B13"].data_type, "s")

    def test_customs_no_freight_or_grand_total(self):
        invoice.build_document(self.data, self.output, "customs")
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        values = [cell.value for row in wb.active for cell in row if cell.value]
        self.assertIn("Subtotal (goods)", values)
        self.assertFalse(any("TOTAL (" in str(value) or value == "Freight" for value in values))

    def test_packing_missing_measurements_blank(self):
        invoice.build_document(example("packing.json"), self.output, "packing")
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active["H13"].value, 0.024)
        self.assertIsNone(wb.active["H14"].value)
        self.assertIsNone(wb.active["F14"].value)

    def test_freight_only_on_commercial_documents(self):
        self.data["freight"] = "10.01"
        self.assertEqual(
            invoice.build_document(self.data, self.output, "proforma")["total"], "152.05"
        )
        for kind in ("customs", "packing"):
            with self.subTest(kind=kind), self.assertRaises(ValueError):
                invoice.build_document(self.data, self.output, kind)

    def test_reject_invalid_or_internal_data_before_output(self):
        for invalid in (True, "NaN", "Infinity", -1, "1e999", "1e-999", None):
            data = copy.deepcopy(self.data)
            data["items"][0]["quantity"] = invalid
            with self.subTest(value=invalid), self.assertRaises(ValueError):
                invoice.build_document(data, self.output)
            self.assertFalse(self.output.exists())
        self.data["items"][0]["supplier_cost"] = 2
        with self.assertRaises(ValueError):
            invoice.build_document(self.data, self.output)

    def test_never_overwrite_existing_output(self):
        self.output.write_bytes(b"keep me")
        with self.assertRaises(FileExistsError):
            invoice.build_document(self.data, self.output)
        self.assertEqual(self.output.read_bytes(), b"keep me")


class FreightTests(unittest.TestCase):
    def test_per_carton_rounding_and_surcharges(self):
        result = freight.estimate(example("freight.json"))
        self.assertEqual(result["total_chargeable_kg"], "13.0")
        self.assertEqual(result["base_charge"], "58.50")
        self.assertEqual(result["total"], "69.35")

    def test_sea_minimum_volume(self):
        result = freight.estimate(example("freight-sea.json"))
        self.assertEqual(result["total_cbm"], "0.048")
        self.assertEqual(result["billable_units"], "1")
        self.assertEqual(result["total"], "80.00")
        self.assertIsNone(result["total_chargeable_kg"])

    def test_minimum_charge_before_surcharges(self):
        data = example("freight.json")
        data["minimum_charge"] = 100
        self.assertEqual(freight.estimate(data)["total"], "115.00")

    def test_volumetric_weight_and_size_warning(self):
        data = example("freight.json")
        data["cartons"][0]["dimensions_cm"] = [100, 50, 50]
        data["max_length_cm"] = 80
        result = freight.estimate(data)
        self.assertEqual(result["total_chargeable_kg"], "100")
        self.assertTrue(any("exceeds" in text for text in result["warnings"]))

    def test_missing_rate_never_means_free(self):
        for rate in (None, 0, True, "NaN", -1, "1e999"):
            data = example("freight.json")
            data["rate"] = rate
            with self.subTest(rate=rate), self.assertRaises(ValueError):
                freight.estimate(data)

    def test_invalid_carton_and_source(self):
        for field, value in (
            ("count", 1.5),
            ("dimensions_cm", [20, 30]),
            ("gross_weight_kg", 0),
            ("dimensions_source", "guessed"),
        ):
            data = example("freight.json")
            data["cartons"][0][field] = value
            with self.subTest(field=field), self.assertRaises(ValueError):
                freight.estimate(data)

    def test_misspelled_fee_is_rejected(self):
        data = example("freight.json")
        data["surchage_percent"] = 20
        with self.assertRaises(ValueError):
            freight.estimate(data)


class TranslationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.xlsx"
        self.map = Path(self.temp.name) / "map.json"
        self.output = Path(self.temp.name) / "translated.xlsx"
        demo.build_demo(self.source)
        glossary = json.loads(
            (ROOT / "skills/translate-rfq/assets/es-zh-demo.json").read_text(encoding="utf-8")
        )
        translate.extract(self.source, self.map, glossary)

    def test_inline_shared_rich_text_and_images_preserved(self):
        source_before = self.source.read_bytes()
        result = translate.apply(self.source, self.map, self.output)
        self.assertEqual(result["images_preserved"], 1)
        self.assertEqual(result["translated_nodes"], 5)
        self.assertEqual(self.source.read_bytes(), source_before)
        with zipfile.ZipFile(self.source) as old, zipfile.ZipFile(self.output) as new:
            for name in old.namelist():
                ET.fromstring(new.read(name)) if name.endswith(".xml") else None
                if name not in result["changed_members"]:
                    self.assertEqual(old.read(name), new.read(name), name)
            self.assertIn(b"<rPr><b/></rPr>", new.read("xl/sharedStrings.xml"))
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        self.assertEqual(wb["RFQ"]["A3"].value, "支架与安装 <演示>")
        self.assertEqual(wb["RFQ"]["B2"].value, "DEMO-A001")
        self.assertEqual(wb["RFQ"]["E2"].value, "=C2*D2")
        self.assertEqual(wb["Notes"]["A1"].value, "询价单")

    def test_namespace_and_phonetic_guides(self):
        xml = b'<s:sst xmlns:s="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><s:si><s:t>Junta</s:t><s:rPh sb="0" eb="5"><s:t>Junta</s:t></s:rPh></s:si></s:sst>'
        result, count = translate.rewrite_xml(xml, {"Junta": "垫片 & <密封>"})
        self.assertEqual(count, 1)
        self.assertIn(b'<s:rPh sb="0" eb="5"><s:t>Junta</s:t></s:rPh>', result)
        self.assertEqual(translate.xml_strings(result), ["垫片 & <密封>"])

    def test_identifiers_are_protected(self):
        self.assertTrue(translate.protected("demo-a001"))
        data = json.loads(self.map.read_text(encoding="utf-8"))
        data["map"]["DEMO-A001"] = "Changed code"
        self.map.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaises(ValueError):
            translate.apply(self.source, self.map, self.output)
        self.assertFalse(self.output.exists())

    def test_invalid_xml_character_rejected(self):
        with self.assertRaises(ET.ParseError):
            translate.rewrite_xml(b"<t>Junta</t>", {"Junta": "bad\x00"})

    def test_map_must_match_source(self):
        with zipfile.ZipFile(self.source, "a") as archive:
            archive.writestr("extra.txt", "changed")
        with self.assertRaises(ValueError):
            translate.apply(self.source, self.map, self.output)
        self.assertFalse(self.output.exists())

    def test_no_overwrite_and_unknown_keys_rejected(self):
        self.output.write_bytes(b"keep me")
        with self.assertRaises(FileExistsError):
            translate.apply(self.source, self.map, self.output)
        self.assertEqual(self.output.read_bytes(), b"keep me")
        data = json.loads(self.map.read_text(encoding="utf-8"))
        data["map"]["Not in workbook"] = "unknown"
        self.map.write_text(json.dumps(data), encoding="utf-8")
        self.output.unlink()
        with self.assertRaises(ValueError):
            translate.apply(self.source, self.map, self.output)


if __name__ == "__main__":
    unittest.main()
